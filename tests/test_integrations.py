"""Regression tests for the vendored/extracted open-source integrations.

These tests deliberately avoid network access and optional OCR models so they can
run in CI on a stock runner. Tests that need an optional dependency skip when it
is unavailable.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.services import pdf_layout, text_normalize
from app.services.extractor import SUPPORTED_EXTENSIONS, extract_paragraphs, iter_docx_text_paragraphs
from app.services.ocr import group_lines
from app.services.pdf_layout import available as pdf_layout_available
from app.services.segment_batch import apply_action, matches_filter, normalize_action, normalize_filter, run_batch
from app.services.segmenter import detect_language, join_lines, split_into_segments, split_sentences

VENDOR = Path(__file__).resolve().parents[1] / "app" / "static" / "vendor"


# --- zhconv-backed Chinese script normalisation ----------------------------


def test_normalize_mode_accepts_known_values_and_falls_back_to_auto():
    assert text_normalize.normalize_mode("TRADITIONAL") == "traditional"
    assert text_normalize.normalize_mode(None) == "auto"
    assert text_normalize.normalize_mode("nonsense") == "auto"


def test_auto_mode_only_converts_for_chinese_target():
    assert text_normalize.effective_mode("auto", "zh") == "simplified"
    assert text_normalize.effective_mode("auto", "en") == "off"
    assert text_normalize.effective_mode("traditional", "en") == "traditional"


def test_off_mode_leaves_text_untouched():
    assert text_normalize.convert_script("這個軟件", "off") == "這個軟件"


@pytest.mark.skipif(not text_normalize.available(), reason="zhconv is not installed")
def test_traditional_and_simplified_conversion():
    traditional = "這個軟件支持繁體字與簡體字轉換。"
    simplified = "这个软件支持繁体字与简体字转换。"
    assert text_normalize.convert_script(traditional, "simplified") == simplified
    # zh-tw output also normalises regional vocabulary (軟件 -> 軟體).
    assert text_normalize.convert_script(simplified, "traditional") == "這個軟體支持繁體字與簡體字轉換。"
    # Round-tripping back to Simplified is stable.
    assert text_normalize.convert_script(
        text_normalize.convert_script(simplified, "traditional"), "simplified"
    ) == simplified


@pytest.mark.skipif(not text_normalize.available(), reason="zhconv is not installed")
def test_normalize_translation_is_a_noop_for_english_targets():
    value = "這個軟件"
    assert text_normalize.normalize_translation(value, "en", "auto") == value


# --- pySBD-backed sentence splitting ---------------------------------------


def test_sentence_split_keeps_abbreviations_and_decimals_intact():
    sentences = split_sentences("The U.S. version 2.1 costs 3.5 dollars, e.g. cheap. It works! See Dr. Smith.")
    assert sentences == [
        "The U.S. version 2.1 costs 3.5 dollars, e.g. cheap.",
        "It works!",
        "See Dr. Smith.",
    ]


def test_sentence_split_handles_chinese_punctuation():
    assert split_sentences("这是第一句。这是第二句！还有第三句？") == ["这是第一句。", "这是第二句！", "还有第三句？"]


def test_long_paragraphs_are_chunked_without_losing_text():
    text = "。".join(["句子" * 20] * 6) + "。"
    chunks = split_into_segments([text], max_chars=60)
    assert all(len(chunk) <= 60 for chunk in chunks)
    assert "".join(chunks) == text


def test_join_lines_omits_spaces_only_across_cjk():
    assert join_lines(["组合驾驶", "辅助系统"]) == "组合驾驶辅助系统"
    assert join_lines(["senior", "professional"]) == "senior professional"
    assert detect_language("组合驾驶辅助系统") == "zh"


# --- DOCX paragraph iteration determinism ----------------------------------


def _build_docx(path: Path, *, merged: bool = False) -> None:
    document = Document()
    document.add_heading("Integration Fixture", level=1)
    document.add_paragraph("The quick brown fox jumps over the lazy dog.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Torque"
    table.cell(1, 1).text = "12 N·m"
    if merged:
        table.cell(1, 0).merge(table.cell(1, 1))
    document.save(path)


def test_iter_docx_text_paragraphs_is_stable_across_reads(tmp_path: Path):
    # `id()` based de-duplication used to return a different paragraph count for
    # the same file because python-docx recycles temporary cell proxies.
    path = tmp_path / "fixture.docx"
    _build_docx(path)
    counts = {len(list(iter_docx_text_paragraphs(Document(path)))) for _ in range(25)}
    assert counts == {6}


def test_table_cells_follow_body_paragraphs(tmp_path: Path):
    path = tmp_path / "fixture.docx"
    _build_docx(path)
    paragraphs = [paragraph.text for paragraph in iter_docx_text_paragraphs(Document(path))]
    assert paragraphs == [
        "Integration Fixture",
        "The quick brown fox jumps over the lazy dog.",
        "Field",
        "Value",
        "Torque",
        "12 N·m",
    ]


def test_merged_cells_are_not_repeated(tmp_path: Path):
    path = tmp_path / "merged.docx"
    _build_docx(path, merged=True)
    paragraphs = [paragraph.text for paragraph in iter_docx_text_paragraphs(Document(path))]
    # A merged cell spans both columns, so its paragraphs must be reported once.
    assert paragraphs.count("Torque") == 1
    assert paragraphs.count("12 N·m") == 1


def test_extractor_accepts_docx_and_rejects_unsupported_suffixes(tmp_path: Path):
    path = tmp_path / "fixture.docx"
    _build_docx(path)
    assert extract_paragraphs(path) == [
        "Integration Fixture",
        "The quick brown fox jumps over the lazy dog.",
        "Field",
        "Value",
        "Torque",
        "12 N·m",
    ]
    assert SUPPORTED_EXTENSIONS == {".doc", ".docx", ".pdf"}
    with pytest.raises(Exception):
        extract_paragraphs(tmp_path / "fixture.txt")


# --- pdfplumber-backed PDF extraction --------------------------------------


def test_pdf_layout_reports_availability_without_raising():
    assert isinstance(pdf_layout_available(), bool)


def test_pdf_mode_names_are_normalised():
    assert pdf_layout.MODE_VALUES == ("plain", "layout")
    assert pdf_layout.normalize_mode("LAYOUT") == "layout"
    assert pdf_layout.normalize_mode(None) == pdf_layout.PLAIN_MODE
    assert pdf_layout.normalize_mode("nonsense") == pdf_layout.PLAIN_MODE


@pytest.mark.skipif(not pdf_layout_available(), reason="pdfplumber is not installed")
def test_two_column_page_is_split_on_the_gutter():
    def word(text, x0, x1, top):
        return {"text": text, "x0": x0, "x1": x1, "top": top}

    # Left column at x 40-240, right column at x 300-500, nothing straddling.
    words = [word(f"L{index}", 40, 240, 10 + index * 12) for index in range(12)]
    words += [word(f"R{index}", 300, 500, 10 + index * 12) for index in range(12)]

    text = pdf_layout._page_text_from_words(_FakePage(words, width=540))
    lines = text.split("\n")
    assert lines[:2] == ["L0", "L1"]
    assert lines[12:14] == ["R0", "R1"]


@pytest.mark.skipif(not pdf_layout_available(), reason="pdfplumber is not installed")
def test_single_column_page_keeps_reading_order():
    def word(text, x0, x1, top):
        return {"text": text, "x0": x0, "x1": x1, "top": top}

    words = [word("first", 40, 200, 10), word("second", 210, 400, 10),
             word("third", 40, 200, 22), word("fourth", 210, 400, 22)]

    assert pdf_layout._page_text_from_words(_FakePage(words, width=540)) == "first second\nthird fourth"


@pytest.mark.skipif(not pdf_layout_available(), reason="pdfplumber is not installed")
def test_cjk_rows_are_joined_without_spaces():
    def word(text, x0, x1, top):
        return {"text": text, "x0": x0, "x1": x1, "top": top}

    words = [word("组合驾驶", 40, 140, 10), word("辅助系统", 140, 240, 10)]

    assert pdf_layout._page_text_from_words(_FakePage(words, width=540)) == "组合驾驶辅助系统"


class _FakePage:
    """Minimal stand-in for a pdfplumber page in the extractor tests."""

    def __init__(self, words: list[dict], width: float = 540):
        self._words = words
        self.width = width

    def extract_words(self, **_kwargs):
        return list(self._words)


@pytest.mark.skipif(not pdf_layout_available(), reason="pdfplumber is not installed")
def test_two_column_pdf_is_read_in_layout_order():
    pdfplumber = pytest.importorskip("pdfplumber")
    path = Path(__file__).parent / "data" / "two-column.pdf"
    if not path.is_file():
        pytest.skip("fixture PDF is not checked in")
    with pdfplumber.open(str(path)) as pdf:
        text = pdf.pages[0].extract_text(layout=True) or ""
    assert text.strip()


# --- RapidOCR line grouping ------------------------------------------------


def test_ocr_grouping_splits_on_vertical_gaps():
    lines = [
        (10.0, 40.0, "First line of paragraph one"),
        (22.0, 40.0, "second line of paragraph one"),
        (80.0, 40.0, "A separate paragraph"),
    ]
    assert group_lines(lines) == ["First line of paragraph one second line of paragraph one", "A separate paragraph"]


# --- vendored front-end assets --------------------------------------------


@pytest.mark.parametrize(
    "name,marker",
    [
        ("purify.min.js", b"DOMPurify"),
        ("diff.min.js", b"diffLines"),
        ("markdown-it.min.js", b"markdown-it"),
        ("markdown-it-task-lists.min.js", b"markdown-it-task-lists"),
    ],
)
def test_vendored_scripts_are_present_and_served_locally(name: str, marker: bytes):
    path = VENDOR / name
    assert path.is_file(), f"{name} must ship with the app (no CDN access at runtime)"
    assert marker in path.read_bytes()[:4096] or marker in path.read_bytes()


@pytest.mark.parametrize(
    "license_name",
    ["purify.LICENSE.txt", "diff.LICENSE.txt", "markdown-it.LICENSE.txt", "markdown-it-task-lists.LICENSE.txt"],
)
def test_vendored_licenses_are_shipped(license_name: str):
    assert (VENDOR / license_name).is_file()


def test_index_html_loads_vendor_scripts_before_feature_modules():
    index = (VENDOR.parent / "index.html").read_text(encoding="utf-8")
    for asset in ("purify.min.js", "diff.min.js", "markdown-it.min.js", "markdown-it-task-lists.min.js"):
        assert f"/vendor/{asset}" in index
    # The feature modules read these globals at definition time, so load order
    # is load-bearing rather than cosmetic.
    assert index.index("/vendor/markdown-it.min.js") < index.index("/scripts/features/markdown.js")
    assert index.index("/vendor/diff.min.js") < index.index("/scripts/features/compare.js")


# The page is served from `app/static/`, and every id app.js looks up with
# `$("#…")` must exist in the markup. A typo here is invisible until the element
# is first touched at runtime, which for a dialog can be a long way from load.
def test_every_dom_id_referenced_by_app_js_exists_in_index():
    static_dir = VENDOR.parent
    html = (static_dir / "index.html").read_text(encoding="utf-8")
    script = (static_dir / "app.js").read_text(encoding="utf-8")
    ids = set(re.findall(r'id=["\']([^"\']+)', html))
    referenced = set(re.findall(r'\$\(["\']#([^"\']+)', script))
    missing = sorted(referenced - ids)
    assert not missing, f"app.js looks up ids that index.html does not define: {missing}"


# Design tokens are the single source of truth for colour and type. Anything
# outside tokens.css must consume them through var(), never a literal.
def test_functional_stylesheets_define_no_colour_literals():
    styles = VENDOR.parent / "styles"
    offenders = {}
    for path in sorted(styles.glob("*.css")):
        if path.name == "tokens.css":
            continue
        text = path.read_text(encoding="utf-8")
        literals = sorted(set(re.findall(r"#[0-9a-fA-F]{3,8}\b", text)))
        literals += sorted(set(re.findall(r"\brgba?\([^)]*\)", text)))
        if literals:
            offenders[path.name] = literals[:6]
    assert not offenders, f"hardcoded colours outside tokens.css: {offenders}"


def test_every_used_token_is_defined_and_vice_versa():
    styles = VENDOR.parent / "styles"
    tokens = (styles / "tokens.css").read_text(encoding="utf-8")
    defined = set(re.findall(r"^\s*--([a-z0-9-]+)\s*:", tokens, re.M))
    used: set[str] = set()
    for path in sorted(styles.glob("*.css")):
        if path.name == "tokens.css":
            continue
        used |= set(re.findall(r"var\(--([a-z0-9-]+)\)", path.read_text(encoding="utf-8")))
    assert not (used - defined), f"used but undefined tokens: {sorted(used - defined)}"
    # The static scales are a palette (docs/DESIGN.md §2.1), so an unused step is
    # allowed; every alias must have a consumer.
    unused = sorted(t for t in defined - used if not t.startswith("static-") and t != "md-scale")
    assert not unused, f"defined but unused tokens: {unused}"


def test_docx_package_is_not_broken_by_the_paragraph_iterator(tmp_path: Path):
    # Guards the export path: the DOCX writer re-reads the copied template and
    # requires exactly the same number of paragraphs as the import pass.
    path = tmp_path / "fixture.docx"
    _build_docx(path)
    imported = list(iter_docx_text_paragraphs(Document(path)))
    copied = tmp_path / "copy.docx"
    copied.write_bytes(path.read_bytes())
    assert len(list(iter_docx_text_paragraphs(Document(copied)))) == len(imported)
    assert zipfile.is_zipfile(copied)


# --- batch segment operations ---------------------------------------------


def _segment(status: str = "empty", translation: str = "", locked: bool = False) -> dict:
    return {"id": f"s-{status}-{translation[:4]}", "translation": translation, "status": status, "locked": locked}


def test_unknown_filter_falls_back_to_all_and_unknown_action_raises():
    assert normalize_filter("nonsense") == "all"
    assert normalize_filter("REVIEWED") == "reviewed"
    with pytest.raises(ValueError):
        normalize_action("delete")


def test_untranslated_filter_covers_blank_and_queued_segments():
    assert matches_filter(_segment("empty"), "untranslated")
    assert matches_filter(_segment("queued"), "untranslated")
    assert not matches_filter(_segment("machine", "译文"), "untranslated")


def test_machine_filter_ignores_edited_and_reviewed_segments():
    assert matches_filter(_segment("machine", "译文"), "machine")
    assert not matches_filter(_segment("edited", "译文"), "machine")
    assert not matches_filter(_segment("reviewed", "译文"), "machine")


def test_run_batch_counts_matches_and_changes():
    segments = [
        _segment("machine", "A"),
        _segment("edited", "B"),
        _segment("reviewed", "C", locked=True),
        _segment("empty"),
    ]
    result = run_batch(segments, "machine", "review")
    assert (result.matched, result.changed) == (1, 1)
    assert segments[0]["status"] == "reviewed"
    assert segments[0]["locked"] is True


def test_review_skips_segments_without_translation():
    segments = [_segment("empty")]
    assert run_batch(segments, "all", "review").changed == 0


def test_batch_actions_leave_locked_segments_alone():
    # The lock is the user's promise a paragraph is settled; bulk review and
    # bulk clear must not silently break it.
    locked = _segment("machine", "译文", locked=True)
    assert apply_action(locked, "review") is False
    assert apply_action(locked, "clear") is False
    assert locked["status"] == "machine"
    assert locked["translation"] == "译文"


def test_unlock_is_the_only_action_that_targets_locked_segments():
    locked = _segment("reviewed", "译文", locked=True)
    assert run_batch([locked], "locked", "unlock").changed == 1
    assert locked["locked"] is False
    assert locked["status"] == "reviewed"


def test_unreview_returns_a_reviewed_segment_to_edited():
    segment = _segment("reviewed", "译文", locked=True)
    assert apply_action(segment, "unreview") is True
    assert segment["status"] == "edited"
    assert segment["locked"] is False


def test_clear_empties_translation_and_status():
    segment = _segment("machine", "译文")
    assert apply_action(segment, "clear") is True
    assert segment["translation"] == ""
    assert segment["status"] == "empty"


def test_batch_is_idempotent():
    segments = [_segment("machine", "A")]
    first = run_batch(segments, "all", "lock")
    second = run_batch(segments, "all", "lock")
    assert (first.matched, first.changed) == (1, 1)
    assert (second.matched, second.changed) == (1, 0)