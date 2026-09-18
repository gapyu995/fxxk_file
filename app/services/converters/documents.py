"""Document conversion: HTML to Markdown and Markdown/DOCX round trips.

HTML conversion uses markdownify (MIT, matthewwithanm/python-markdownify). DOCX
writing stays on python-docx, which the translation pipeline already depends on,
so only the HTML direction can be missing at runtime.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt

try:  # pragma: no cover - exercised through the public helpers below
    from markdownify import markdownify as _markdownify
except Exception:  # pragma: no cover - optional dependency guard
    _markdownify = None


class DocumentConversionError(RuntimeError):
    pass


def html_available() -> bool:
    return _markdownify is not None


def unavailable_reason() -> str:
    return "未安装 markdownify（pip install -r requirements-convert.txt）。"


def html_to_markdown(html: str) -> str:
    """Convert HTML to Markdown, dropping scripts and styles."""
    if _markdownify is None:
        raise DocumentConversionError(unavailable_reason())
    if not str(html or "").strip():
        raise DocumentConversionError("HTML 内容为空。")
    cleaned = re.sub(r"(?is)<(script|style)\b.*?</\1>", "", str(html))
    markdown = _markdownify(
        cleaned,
        heading_style="ATX",
        bullets="-",
        strip=["script", "style"],
    )
    return re.sub(r"\n{3,}", "\n\n", markdown).strip() + "\n"


def docx_to_markdown(data: bytes) -> str:
    """Convert a DOCX body (paragraphs and tables, in document order) to Markdown."""
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentConversionError(f"DOCX 无法读取：{exc}") from exc
    lines: list[str] = []
    for block in _iter_body_blocks(document):
        # A paragraph yields one Markdown string, a table yields the table
        # object; extending a string here would split it into characters.
        lines.append(block if isinstance(block, str) else _table_to_markdown(block))
    text = "\n\n".join(part for part in (chunk.strip("\n") for chunk in lines) if part)
    if not text.strip():
        raise DocumentConversionError("DOCX 中没有可转换的文字。")
    return text.strip() + "\n"


def markdown_to_docx(markdown: str) -> bytes:
    """Convert a Markdown subset (headings, lists, tables, quotes, code) to DOCX."""
    document = Document()
    lines = str(markdown or "").replace("\r\n", "\n").split("\n")
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph("\n".join(code))
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            document.add_heading(_inline_text(heading.group(2)), level=len(heading.group(1)))
            index += 1
            continue
        if _is_table_start(lines, index):
            header, rows, index = _read_table(lines, index)
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for cell, value in zip(table.rows[0].cells, header):
                cell.text = value
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
            continue
        if stripped.startswith(">"):
            quote: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote.append(lines[index].strip().lstrip(">").strip())
                index += 1
            document.add_paragraph(_inline_text(" ".join(quote)), style="Intense Quote")
            continue
        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        ordered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if bullet or ordered:
            style = "List Number" if ordered else "List Bullet"
            while index < len(lines):
                candidate = lines[index].strip()
                match = re.match(r"^(?:[-*+]\s+|\d+[.)]\s+)(.*)$", candidate) if ordered else re.match(r"^[-*+]\s+(.*)$", candidate)
                if not match:
                    break
                document.add_paragraph(_inline_text(match.group(1)), style=style)
                index += 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            document.add_paragraph("—" * 20)
            index += 1
            continue
        document.add_paragraph(_inline_text(stripped))
        index += 1
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _iter_body_blocks(document):
    """Yield paragraph Markdown strings and table objects in body order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            markdown = _paragraph_to_markdown(paragraph)
            if markdown is not None:
                yield markdown
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _paragraph_to_markdown(paragraph) -> str | None:
    text = paragraph.text.strip()
    if not text:
        return None
    style = (paragraph.style.name or "").lower()
    heading = re.match(r"heading (\d)", style)
    if heading:
        return ["#", "##", "###", "####", "#####", "######"][min(5, int(heading.group(1)) - 1)] + " " + text
    if "list bullet" in style:
        return "- " + text
    if "list number" in style:
        return "1. " + text
    if "quote" in style:
        return "> " + text
    return text


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(padded[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return "\n".join(lines)


def _inline_text(value: str) -> str:
    """Keep the text, drop the Markdown emphasis markers python-docx cannot show."""
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", value)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", value and text)
    text = re.sub(r"(\*\*|__|\*|_|`)", "", text)
    return text


def _is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    if "|" not in lines[index]:
        return False
    separator = lines[index + 1].strip().strip("|").strip()
    if not separator:
        return False
    cells = [cell.strip() for cell in separator.split("|")]
    return all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell)


def _read_table(lines: list[str], index: int) -> tuple[list[str], list[list[str]], int]:
    header = _split_row(lines[index])
    index += 2
    rows: list[list[str]] = []
    while index < len(lines) and "|" in lines[index] and lines[index].strip():
        row = _split_row(lines[index])
        rows.append(row + [""] * (len(header) - len(row)))
        index += 1
    return header, rows, index


def _split_row(line: str) -> list[str]:
    value = line.strip().strip("|")
    return [cell.strip().replace("\\|", "|") for cell in value.split("|")]
